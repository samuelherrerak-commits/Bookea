# Tarea 3 de Estadísticas y Probabilidades: libro de Excel con fórmulas vivas y gráficos nativos.
# Se usan las funciones clásicas (BINOMDIST, NORMSDIST, NORMSINV, NORMDIST, NORMINV) que Excel y LibreOffice entienden sin prefijo.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, AreaChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

F='Arial'
f_norm=Font(name=F,size=10); f_b=Font(name=F,size=10,bold=True); f_in=Font(name=F,size=10,color='0000FF',bold=True)
f_h=Font(name=F,size=10,bold=True,color='FFFFFF'); f_t=Font(name=F,size=14,bold=True); f_s=Font(name=F,size=10,italic=True,color='5B5B57')
fill_h=PatternFill('solid',fgColor='0F0F0E'); fill_in=PatternFill('solid',fgColor='FFFF00'); fill_res=PatternFill('solid',fgColor='F4F4F2'); fill_sec=PatternFill('solid',fgColor='E3E3DF')
thin=Side(style='thin',color='BFBFBF'); box=Border(top=thin,bottom=thin,left=thin,right=thin)
wrap=Alignment(wrap_text=True,vertical='top'); ctr=Alignment(horizontal='center',vertical='center',wrap_text=True)

wb=Workbook()
def hoja(nombre,primero=False):
    ws=wb.active if primero else wb.create_sheet()
    ws.title=nombre; ws.sheet_view.showGridLines=False; return ws
def celda(ws,ref,v,font=None,fmt=None,fill=None,al=None,border=True):
    c=ws[ref]; c.value=v; c.font=font or f_norm
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    if al: c.alignment=al
    if border: c.border=box
    return c
def encabezados(ws,fila,cols,desde=1):
    for i,t in enumerate(cols):
        celda(ws,f'{L(desde+i)}{fila}',t,f_h,fill=fill_h,al=ctr)
def titulo(ws,t,s):
    ws['A1']=t; ws['A1'].font=f_t; ws['A2']=s; ws['A2'].font=f_s

R=hoja('Resumen',True); P=hoja('Parámetros'); B=hoja('Binomial'); N=hoja('Normal'); T=hoja('Tabla Z'); S=hoja('Sensibilidad')

# ---------------- Parámetros ----------------
titulo(P,'Parámetros del problema','Los datos de entrada (letra azul con fondo amarillo) vienen de la consigna de la Tarea 3; todas las demás hojas se recalculan si se cambian.')
encabezados(P,4,['Parámetro','Valor','Descripción y fuente'])
par=[('n (unidades por lote)',200,'0','Consigna: en un lote típico se producen 200 unidades por día.'),
 ('p (probabilidad de defecto)',0.03,'0.00','Consigna: tasa promedio histórica de defectos de 3 % por unidad.'),
 ('q = 1 − p','=1-B6','0.00','Probabilidad de que una unidad salga sin defecto.'),
 ('μ del tiempo de procesamiento (min)',45,'0','Consigna: media de 45 minutos.'),
 ('σ del tiempo de procesamiento (min)',8,'0','Consigna: desviación estándar de 8 minutos.'),
 ('k para P(X = k)',5,'0','Consigna, Parte A.4: exactamente 5 defectos.'),
 ('k para P(X ≥ k)',6,'0','Consigna, Parte A.5: 6 o más defectos.'),
 ('k para P(X ≤ k)',3,'0','Consigna, Parte A.6: como máximo 3 defectos.'),
 ('x₁ (min)',37,'0','Consigna, Parte B: límite inferior del intervalo de tiempo.'),
 ('x₂ (min)',53,'0','Consigna, Parte B: límite superior del intervalo de tiempo.'),
 ('Percentil bajo (10 % más rápido)',0.10,'0%','Consigna, pregunta 4: el 10 % más rápido.'),
 ('Percentil alto',0.90,'0%','Consigna, Parte B.7: percentil 90.')]
for i,(a,v,fm,d) in enumerate(par):
    r=5+i; celda(P,f'A{r}',a,f_b); entrada=not (isinstance(v,str) and v.startswith('='))
    celda(P,f'B{r}',v,f_in if entrada else f_norm,fm,fill_in if entrada else None,Alignment(horizontal='center')); celda(P,f'C{r}',d,al=wrap)
P['A18']='Convención de colores: azul sobre amarillo = dato de entrada; negro = fórmula. Las probabilidades se muestran con 4 decimales.'; P['A18'].font=f_s
P.column_dimensions['A'].width=40; P.column_dimensions['B'].width=14; P.column_dimensions['C'].width=80
pN,pp,pq,pM,pS,kE,kG,kM,x1,x2,pl,ph=[f"Parámetros!$B${5+i}" for i in range(12)]

# ---------------- Binomial ----------------
titulo(B,'Distribución binomial: defectos en un lote de 200 unidades','P(X = k) = C(n, k) · p^k · (1 − p)^(n − k). Parámetros en la hoja «Parámetros».')
encabezados(B,4,['Medida','Fórmula','Resultado','Interpretación'])
filas=[
 ('μ: valor esperado de defectos','μ = n × p',f'={pN}*{pp}','0.0000','Defectos que se esperan, en promedio, por lote de 200 unidades.'),
 ('σ²: varianza','σ² = n × p × (1 − p)',f'={pN}*{pp}*{pq}','0.0000','Dispersión de los defectos alrededor de la media, en unidades al cuadrado.'),
 ('σ: desviación estándar','σ = √(n × p × (1 − p))','=SQRT(C6)','0.0000','Variación típica del número de defectos por lote.'),
 (f'="C(n, k) con k = "&{kE}','C(n, k) = n! / (k!·(n − k)!)',f'=COMBIN({pN},{kE})','#,##0','Formas de escoger cuáles unidades son las defectuosas.'),
 (f'="p^k con k = "&{kE}','p^k',f'={pp}^{kE}','0.000E+00','Probabilidad de que las k unidades escogidas sean defectuosas.'),
 (f'="(1 − p)^(n − k) con k = "&{kE}','(1 − p)^(n − k)',f'={pq}^({pN}-{kE})','0.000E+00','Probabilidad de que las demás unidades salgan buenas.'),
 (f'="P(X = "&{kE}&")"','C(n, k) · p^k · (1 − p)^(n − k)','=C8*C9*C10','0.0000','Probabilidad de exactamente 5 defectos en el lote.'),
 (f'="P(X = "&{kE}&") con BINOMDIST (verificación)"','BINOMDIST(k; n; p; FALSO)',f'=BINOMDIST({kE},{pN},{pp},FALSE)','0.0000','Debe coincidir con el cálculo manual de la fila anterior.'),
 (f'="P(X ≤ "&({kG}-1)&")"','Suma de P(X = k) desde 0 hasta 5',f'=BINOMDIST({kG}-1,{pN},{pp},TRUE)','0.0000','Probabilidad acumulada de tener hasta 5 defectos.'),
 (f'="P(X ≥ "&{kG}&") = 1 − P(X ≤ "&({kG}-1)&")"','1 − P(X ≤ 5)','=1-C13','0.0000','Probabilidad de tener 6 o más defectos en el lote.'),
 (f'="P(X ≤ "&{kM}&")"','BINOMDIST(3; n; p; VERDADERO)',f'=BINOMDIST({kM},{pN},{pp},TRUE)','0.0000','Probabilidad de un lote con a lo sumo 3 defectos.'),
 ('Rango usual de defectos: μ − 2σ','μ − 2σ','=C5-2*C7','0.00','Límite inferior del intervalo que cubre cerca del 95 % de los lotes.'),
 ('Rango usual de defectos: μ + 2σ','μ + 2σ','=C5+2*C7','0.00','Límite superior del intervalo que cubre cerca del 95 % de los lotes.'),
 ('P(X ≥ 10): lote crítico','1 − P(X ≤ 9)',f'=1-BINOMDIST(9,{pN},{pp},TRUE)','0.0000','Riesgo de un lote con 10 o más defectos (umbral de alerta propuesto).')]
for i,(a,fo,fx,fm,it) in enumerate(filas):
    r=5+i; celda(B,f'A{r}',a,f_b,al=wrap); celda(B,f'B{r}',fo,al=wrap); celda(B,f'C{r}',fx,f_norm,fm,fill_res,Alignment(horizontal='center',vertical='center')); celda(B,f'D{r}',it,al=wrap)
    for cc in 'EFGH': B[f'{cc}{r}'].border=box
    B.merge_cells(f'D{r}:H{r}'); B.row_dimensions[r].height=28
for cc in 'EFGH': B[f'{cc}4'].fill=fill_h
B.merge_cells('D4:H4')
# tabla k = 0..12
r0=21
B[f'A{r0-1}']='Tabla de probabilidades para k = 0 a 12 defectos'; B[f'A{r0-1}'].font=f_b
encabezados(B,r0,['k','C(n, k)','p^k','(1 − p)^(n − k)','P(X = k) manual','P(X = k) BINOMDIST','P(X ≤ k)','P(X ≥ k)','Serie: P(X ≤ 3)','Serie: k = 4','Serie: k = 5','Serie: k = μ','Serie: k ≥ 6, k ≠ μ'])
for k in range(13):
    r=r0+1+k
    celda(B,f'A{r}',k,f_b,'0',al=Alignment(horizontal='center'))
    celda(B,f'B{r}',f'=COMBIN({pN},A{r})',fmt='0.000E+00'); celda(B,f'C{r}',f'={pp}^A{r}',fmt='0.000E+00'); celda(B,f'D{r}',f'={pq}^({pN}-A{r})',fmt='0.0000')
    celda(B,f'E{r}',f'=B{r}*C{r}*D{r}',fmt='0.0000',fill=fill_res); celda(B,f'F{r}',f'=BINOMDIST(A{r},{pN},{pp},FALSE)',fmt='0.0000')
    celda(B,f'G{r}',f'=BINOMDIST(A{r},{pN},{pp},TRUE)',fmt='0.0000'); celda(B,f'H{r}',f'=IF(A{r}=0,1,1-BINOMDIST(A{r}-1,{pN},{pp},TRUE))',fmt='0.0000')
    celda(B,f'I{r}',f'=IF(A{r}<={kM},E{r},0)',fmt='0.0000'); celda(B,f'J{r}',f'=IF(AND(A{r}>{kM},A{r}<{kE}),E{r},0)',fmt='0.0000')
    celda(B,f'K{r}',f'=IF(A{r}={kE},E{r},0)',fmt='0.0000'); celda(B,f'L{r}',f'=IF(A{r}=$C$5,E{r},0)',fmt='0.0000')
    celda(B,f'M{r}',f'=IF(AND(A{r}>={kG},A{r}<>$C$5),E{r},0)',fmt='0.0000')
rf=r0+13
celda(B,f'A{rf+1}','Σ k = 0..12',f_b); celda(B,f'E{rf+1}',f'=SUM(E{r0+1}:E{rf})',f_b,'0.0000',fill_res)
B[f'A{rf+2}']='Nota: las columnas «Serie» solo separan las barras por color para el gráfico (valor cero = sin barra). Σ queda en 0,9922 porque faltan las colas k ≥ 13.'; B[f'A{rf+2}'].font=f_s
for c,w in zip('ABCDEFGHIJKLM',[44,34,16,60,16,18,12,12,14,12,12,12,16]): B.column_dimensions[c].width=w
B.column_dimensions['D'].width=16
# gráfico binomial
ch=BarChart(); ch.type='col'; ch.grouping='clustered'; ch.overlap=100; ch.gapWidth=30
ch.title='Distribución binomial de defectos por lote (n = 200, p = 0,03)'; ch.y_axis.title='Probabilidad P(X = k)'; ch.x_axis.title='Número de unidades defectuosas (k)'
cats=Reference(B,min_col=1,min_row=r0+1,max_row=rf)
nombres=['P(X ≤ 3)','k = 4','k = 5: P(X = 5)','k = 6 = μ (valor esperado)','k ≥ 7 (con k = 6: P(X ≥ 6))']; colores=['2A78D6','A3A39E','EB6834','0F7A5C','1BAF7A']
for i,(nm,co) in enumerate(zip(nombres,colores)):
    s=Series(Reference(B,min_col=9+i,min_row=r0+1,max_row=rf),title=nm); s.graphicalProperties.solidFill=co; s.graphicalProperties.line.solidFill=co; ch.series.append(s)
ch.set_categories(cats); ch.y_axis.number_format='0.00'; ch.y_axis.scaling.min=0; ch.y_axis.scaling.max=0.18; ch.y_axis.majorGridlines=ch.y_axis.majorGridlines
ch.x_axis.delete=False; ch.y_axis.delete=False; ch.legend.position='b'; ch.height=10; ch.width=22
B.add_chart(ch,f'A{rf+4}')

# ---------------- Tabla Z ----------------
titulo(T,'Tabla de la distribución normal estándar: Φ(z) = P(Z ≤ z), para z ≥ 0','Cada celda es =REDONDEAR(DISTR.NORM.ESTAND(z); 4), con z = encabezado de fila + encabezado de columna. Para z negativo: Φ(−z) = 1 − Φ(z).')
celda(T,'A4','z',f_h,fill=fill_h,al=ctr)
for j in range(10): celda(T,f'{L(2+j)}4',round(j/100,2),f_h,'0.00',fill_h,ctr)
for i in range(35):
    r=5+i; celda(T,f'A{r}',round(i/10,1),f_h,'0.0',fill_h,ctr)
    for j in range(10): celda(T,f'{L(2+j)}{r}',f'=ROUND(NORMSDIST($A{r}+{L(2+j)}$4),4)',fmt='0.0000',al=Alignment(horizontal='center'))
T.column_dimensions['A'].width=8
for j in range(10): T.column_dimensions[L(2+j)].width=9
T.freeze_panes='B5'

# ---------------- Normal ----------------
titulo(N,'Distribución normal: tiempos de procesamiento','z = (x − μ)/σ. Se muestra el valor exacto (Excel) y el valor que resulta de leer la tabla Z con z redondeado a dos decimales.')
celda(N,'A4','μ (min)',f_b); celda(N,'B4',f'={pM}',fmt='0.00',al=Alignment(horizontal='center')); celda(N,'C4','σ (min)',f_b); celda(N,'D4',f'={pS}',fmt='0.00',al=Alignment(horizontal='center'))
encabezados(N,6,['N.º','Cálculo','Fórmula','Resultado exacto (Excel)','Φ(|z|) leído en tabla Z','Resultado con tabla Z','Interpretación'])
tab=lambda z: f"INDEX('Tabla Z'!$B$5:$K$39,INT(ROUND(ABS({z})*100,0)/10)+1,MOD(ROUND(ABS({z})*100,0),10)+1)"
items=[ # fila: (n, calculo, formula, exacto, fmt, tabla, con_tabla, fmt2, interpretacion)
 (1,f'="z para "&{x1}&" min"','z = (x − μ)/σ','=(%s-$B$4)/$D$4'%x1,'0.0000','='+tab('D7'),'=ROUND(D7,2)','0.00','37 min está 1 desviación estándar por debajo de la media.'),
 (2,f'="z para "&{x2}&" min"','z = (x − μ)/σ','=(%s-$B$4)/$D$4'%x2,'0.0000','='+tab('D8'),'=ROUND(D8,2)','0.00','53 min está 1 desviación estándar por encima de la media.'),
 (3,f'="z del percentil "&TEXT({pl},"0%")','z tal que Φ(z) = 0,10','=NORMSINV(%s)'%pl,'0.0000','='+tab('D9'),'=ROUND(D9,2)','0.00','El 10 % más rápido queda 1,28 desviaciones por debajo de la media.'),
 (4,f'="P(X < "&{x1}&")"','Φ(z₁)','=NORMSDIST(D7)','0.0000',None,'=1-E7','0.0000','Probabilidad de que una unidad tarde menos de 37 min.'),
 (5,f'="P(X > "&{x2}&")"','1 − Φ(z₂)','=1-NORMSDIST(D8)','0.0000',None,'=1-E8','0.0000','Probabilidad de que una unidad tarde más de 53 min.'),
 (6,f'="P("&{x1}&" < X < "&{x2}&")"','Φ(z₂) − Φ(z₁)','=NORMSDIST(D8)-NORMSDIST(D7)','0.0000',None,'=E8-(1-E7)','0.0000','Probabilidad de que una unidad tarde entre 37 y 53 min.'),
 (7,f'="Tiempo del percentil "&TEXT({pl},"0%")&" (min)"','x = μ + z·σ','=$B$4+D9*$D$4','0.00',None,'=$B$4+F9*$D$4','0.00','Tiempo que separa el 10 % más rápido del resto.'),
 (8,f'="z del percentil "&TEXT({ph},"0%")','z tal que Φ(z) = 0,90','=NORMSINV(%s)'%ph,'0.0000','='+tab('D14'),'=ROUND(D14,2)','0.00','Simétrico al percentil 10: 1,28 desviaciones por encima de la media.'),
 (9,f'="Tiempo del percentil "&TEXT({ph},"0%")&" (min)"','x = μ + z·σ','=NORMINV(%s,$B$4,$D$4)'%ph,'0.00',None,'=$B$4+F14*$D$4','0.00','El 90 % de las unidades se procesa en este tiempo o menos.')]
for i,(n_,cal,fo,fx,fm,tb,ft,fm2,it) in enumerate(items):
    r=7+i; celda(N,f'A{r}',n_,f_b,al=Alignment(horizontal='center')); celda(N,f'B{r}',cal,f_b,al=wrap); celda(N,f'C{r}',fo,al=wrap)
    celda(N,f'D{r}',fx,f_norm,fm,fill_res,Alignment(horizontal='center')); celda(N,f'E{r}',tb,f_norm,'0.0000',al=Alignment(horizontal='center'))
    celda(N,f'F{r}',ft,f_norm,fm2,fill_res,Alignment(horizontal='center')); celda(N,f'G{r}',it,al=wrap)
# regla empírica
N['A18']='Regla empírica: intervalos de desviación estándar'; N['A18'].font=f_b
encabezados(N,19,['k','Intervalo μ ± kσ','Límite inferior (min)','Límite superior (min)','P(μ − kσ < X < μ + kσ)'])
for k in (1,2,3):
    r=19+k; celda(N,f'A{r}',k,f_b,'0',al=Alignment(horizontal='center')); celda(N,f'B{r}',f'="μ ± "&A{r}&"σ"')
    celda(N,f'C{r}',f'=$B$4-A{r}*$D$4',fmt='0.00'); celda(N,f'D{r}',f'=$B$4+A{r}*$D$4',fmt='0.00'); celda(N,f'E{r}',f'=NORMSDIST(A{r})-NORMSDIST(-A{r})',f_norm,'0.0000',fill_res)
# datos de la curva
c0=27
N[f'A{c0-1}']='Datos para el gráfico: densidad f(x) = NORMDIST(x; μ; σ; FALSO)'; N[f'A{c0-1}'].font=f_b
encabezados(N,c0,['x (min)','z','f(x)','Área x ≤ 37','Área 37 ≤ x ≤ 53','Área x ≥ 53'])
for i,x in enumerate(range(21,70)):
    r=c0+1+i
    celda(N,f'A{r}',x,f_b,'0',al=Alignment(horizontal='center')); celda(N,f'B{r}',f'=(A{r}-$B$4)/$D$4',fmt='0.00'); celda(N,f'C{r}',f'=NORMDIST(A{r},$B$4,$D$4,FALSE)',fmt='0.0000')
    celda(N,f'D{r}',f'=IF(A{r}<={x1},C{r},0)',fmt='0.0000'); celda(N,f'E{r}',f'=IF(AND(A{r}>={x1},A{r}<={x2}),C{r},0)',fmt='0.0000'); celda(N,f'F{r}',f'=IF(A{r}>={x2},C{r},0)',fmt='0.0000')
cf=c0+49
N[f'A{cf+2}']='Nota: las líneas verticales del gráfico marcan la media y los límites μ ± 1σ, 2σ y 3σ (21, 29, 37, 45, 53, 61 y 69 min).'; N[f'A{cf+2}'].font=f_s
for c,w in zip('ABCDEFG',[26,32,26,22,22,22,52]): N.column_dimensions[c].width=w
N.row_dimensions[6].height=32
ar=AreaChart(); ar.grouping='standard'; ar.title='Distribución normal de los tiempos de procesamiento (μ = 45 min, σ = 8 min)'
for col,nm,co in ((4,'x ≤ 37: P = 0,1587','EB6834'),(5,'37 ≤ x ≤ 53: P = 0,6827','2A78D6'),(6,'x ≥ 53: P = 0,1587','1BAF7A')):
    s=Series(Reference(N,min_col=col,min_row=c0+1,max_row=cf),title=nm); s.graphicalProperties.solidFill=co; s.graphicalProperties.line.solidFill=co; ar.series.append(s)
ar.set_categories(Reference(N,min_col=1,min_row=c0+1,max_row=cf))
from openpyxl.chart.axis import ChartLines
ar.y_axis.title='Densidad f(x)'; ar.x_axis.title='Tiempo de procesamiento (min)'; ar.y_axis.number_format='0.00'; ar.x_axis.tickLblSkip=8; ar.x_axis.tickMarkSkip=8; ar.x_axis.majorGridlines=ChartLines(); ar.x_axis.majorGridlines.spPr=None; ar.x_axis.delete=False; ar.y_axis.delete=False
ar.legend.position='b'; ar.height=10; ar.width=22
N.add_chart(ar,'I6')

# ---------------- Sensibilidad ----------------
titulo(S,'Sensibilidad: efecto de la tasa de defectos en el riesgo del lote','Mismo lote (n = 200) con distintas tasas de defecto p; el escenario actual es p = 3 %. Las tasas de la columna A son supuestos de análisis, no datos de la empresa.')
encabezados(S,4,['Tasa de defecto p','μ = n·p','σ','P(X ≤ 3)','P(X ≥ 6)','P(X ≥ 10)','Defectos esperados frente a p = 3 %'])
for i,pv in enumerate((0.01,0.02,0.03,0.04,0.05)):
    r=5+i; celda(S,f'A{r}',pv,f_in,'0%',fill_in,Alignment(horizontal='center'))
    celda(S,f'B{r}',f'={pN}*A{r}',fmt='0.00'); celda(S,f'C{r}',f'=SQRT({pN}*A{r}*(1-A{r}))',fmt='0.00')
    celda(S,f'D{r}',f'=BINOMDIST({kM},{pN},A{r},TRUE)',fmt='0.0000'); celda(S,f'E{r}',f'=1-BINOMDIST({kG}-1,{pN},A{r},TRUE)',fmt='0.0000'); celda(S,f'F{r}',f'=1-BINOMDIST(9,{pN},A{r},TRUE)',fmt='0.0000')
    celda(S,f'G{r}',f'=B{r}/$B$7',fmt='0%')
S['A11']='Lectura: bajar la tasa de defecto de 3 % a 2 % reduce un tercio los defectos esperados y la probabilidad de 6 o más defectos baja de más de la mitad a cerca de un quinto de los lotes.'; S['A11'].font=f_s
for c,w in zip('ABCDEFG',[18,12,10,12,12,12,36]): S.column_dimensions[c].width=w

# ---------------- Resumen ----------------
titulo(R,'Tarea 3: Análisis de probabilidad y dispersión','Planta manufacturera de artículos electrónicos (Los Andes, Mérida). Respuestas al director ejecutivo; todos los valores vienen de las hojas de cálculo.')
encabezados(R,4,['N.º','Pregunta del director ejecutivo','Resultado','Detalle'])
res=[(1,'¿Cuántas unidades defectuosas se espera en un lote de 200 unidades?','=Binomial!C5','0.00','μ = n × p = 200 × 0,03; desviación estándar σ = 2,41 unidades.'),
 (2,'¿Cuál es la probabilidad de tener exactamente 5 unidades defectuosas?','=Binomial!C11','0.0000','P(X = 5) con la fórmula binomial; equivale a 16,22 % de los lotes.'),
 (3,'¿Cuál es la probabilidad de tener 6 o más unidades defectuosas?','=Binomial!C14','0.0000','P(X ≥ 6) = 1 − P(X ≤ 5); más de la mitad de los lotes.'),
 (4,'¿Qué tiempo de procesamiento separa el 10 % más rápido del resto? (min)','=Normal!D13','0.00','Percentil 10 = μ + z·σ = 45 − 1,2816 × 8.'),
 (5,'¿Cuál es la probabilidad de que una unidad tarde entre 37 y 53 minutos?','=Normal!D12','0.0000','P(−1 < Z < 1): regla empírica de 68 %.'),
 (6,'¿Cuáles son las recomendaciones para mejorar la calidad y eficiencia?','Ver documento Word','','Sección «Recomendaciones y conclusiones» del informe.')]
for i,(n_,q,fx,fm,d) in enumerate(res):
    r=5+i; celda(R,f'A{r}',n_,f_b,al=Alignment(horizontal='center',vertical='center')); celda(R,f'B{r}',q,al=wrap); celda(R,f'C{r}',fx,f_b,fm,fill_res,ctr); celda(R,f'D{r}',d,al=wrap); R.row_dimensions[r].height=32
R['A12']='Otros resultados clave'; R['A12'].font=f_b
encabezados(R,13,['','Medida','Resultado','Fuente'])
otros=[('Varianza σ² del número de defectos','=Binomial!C6','0.0000','Hoja Binomial'),('P(X ≤ 3): a lo sumo 3 defectos','=Binomial!C15','0.0000','Hoja Binomial'),
 ('z de 37 min','=Normal!D7','0.00','Hoja Normal'),('z de 53 min','=Normal!D8','0.00','Hoja Normal'),('z del percentil 10','=Normal!D9','0.0000','Hoja Normal'),
 ('P(tiempo < 37 min)','=Normal!D10','0.0000','Hoja Normal'),('P(tiempo > 53 min)','=Normal!D11','0.0000','Hoja Normal'),('Tiempo del percentil 90 (min)','=Normal!D15','0.00','Hoja Normal')]
for i,(m,fx,fm,fu) in enumerate(otros):
    r=14+i; celda(R,f'A{r}',None); celda(R,f'B{r}',m); celda(R,f'C{r}',fx,f_norm,fm,fill_res,Alignment(horizontal='center')); celda(R,f'D{r}',fu)
R['A23']='Hojas del libro: Parámetros (datos de la consigna) · Binomial · Normal · Tabla Z · Sensibilidad. Letra azul sobre amarillo = dato editable; negro = fórmula.'; R['A23'].font=f_s
for c,w in zip('ABCD',[8,62,16,66]): R.column_dimensions[c].width=w
for ws in (R,P,B,N,T,S): ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
B.page_setup.fitToHeight=1
wb.save('../Tarea_3_Estadisticas_y_Probabilidades.xlsx'); print('ok')
